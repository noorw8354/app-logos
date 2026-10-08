"""
Build a self-hosted logo set for the Tableau dashboard.

What it does
1. Reads app_logo_mapping.xlsx (columns: APP, Domain, Logo URL).
2. Downloads each logo from the link currently in the sheet.
3. Saves it as a 512x512 transparent PNG with padding -> ./logos/<app-name>.png
4. Writes app_logo_mapping_hosted.xlsx where every Logo URL points to your GitHub repo
   through cdn.jsdelivr.net (the host that already works in your published dashboard).

Run on your Mac (needs internet):
    pip3 install pillow openpyxl requests
    python3 build_logo_hosting.py
"""
import os, re, sys, io
import requests
from PIL import Image
from openpyxl import load_workbook

# ---------- EDIT THESE THREE LINES ----------
GITHUB_USER = "noorw8354"
GITHUB_REPO = "app-logos"          # public repo, e.g. app-logos
BRANCH      = "main"
# --------------------------------------------

INPUT_XLSX  = "app_logo_mapping.xlsx"
OUTPUT_XLSX = "app_logo_mapping_hosted.xlsx"
LOGO_DIR    = "logos"
SIZE        = 512      # final image is SIZE x SIZE
PAD         = 0.12     # transparent margin around the logo (12% each side) -> clear gaps between rows
HEADERS     = {"User-Agent": "Mozilla/5.0"}

CDN = "https://cdn.jsdelivr.net/gh/{u}/{r}@{b}/{d}/{f}"


def slug(name):
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")


def make_square(img):
    img = img.convert("RGBA")
    inner = int(SIZE * (1 - 2 * PAD))
    img.thumbnail((inner, inner), Image.LANCZOS)
    # upscale small icons so they fill the inner box
    if max(img.size) < inner:
        scale = inner / max(img.size)
        img = img.resize((round(img.width * scale), round(img.height * scale)), Image.LANCZOS)
    canvas = Image.new("RGBA", (SIZE, SIZE), (255, 255, 255, 0))
    canvas.paste(img, ((SIZE - img.width) // 2, (SIZE - img.height) // 2), img)
    return canvas


def main():
    if GITHUB_USER.startswith("YOUR_") or GITHUB_REPO.startswith("YOUR_"):
        sys.exit("Edit GITHUB_USER and GITHUB_REPO at the top of this file first.")
    os.makedirs(LOGO_DIR, exist_ok=True)
    wb = load_workbook(INPUT_XLSX)
    ws = wb.active
    ok, failed, skipped = [], [], []
    for r in range(2, ws.max_row + 1):
        app, url = ws.cell(r, 1).value, ws.cell(r, 3).value
        if not app or not url:
            skipped.append(app)
            continue
        fname = slug(app) + ".png"
        try:
            resp = requests.get(url, headers=HEADERS, timeout=30)
            resp.raise_for_status()
            img = Image.open(io.BytesIO(resp.content))
            make_square(img).save(os.path.join(LOGO_DIR, fname), "PNG")
            ws.cell(r, 3).value = CDN.format(u=GITHUB_USER, r=GITHUB_REPO, b=BRANCH, d=LOGO_DIR, f=fname)
            ok.append(app)
        except Exception as e:
            failed.append((app, str(e)[:80]))
            print(f"FAILED  {app}: {e}")
    wb.save(OUTPUT_XLSX)
    print(f"\nDone. {len(ok)} logos saved in ./{LOGO_DIR}, {len(failed)} failed, {len(skipped)} rows had no link.")
    if failed:
        print("Failed apps keep their old link in the new sheet. Fix those links and run again:")
        for a, e in failed:
            print("  -", a, "|", e)
    print(f"\nNext: upload the '{LOGO_DIR}' folder to https://github.com/{GITHUB_USER}/{GITHUB_REPO} "
          f"and use {OUTPUT_XLSX} in Tableau.")


if __name__ == "__main__":
    main()
